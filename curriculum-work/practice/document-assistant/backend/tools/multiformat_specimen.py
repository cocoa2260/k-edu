"""python -m tools.multiformat_specimen --output NEW_DIR

Synthetic DOCX/TXT inputs. Rendering is a separate mandatory authoring check.
"""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor

ROWS=[['구분','기한','부담금'],['미사용','7일','0원'],['사용 이력','개별 검토','15000원']]


def build(output):
    output=Path(output);output.mkdir(parents=True,exist_ok=False)
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.left_margin=sec.right_margin=Inches(.8);sec.top_margin=sec.bottom_margin=Inches(.8)
    for name in ('Normal','Title','Heading 1'):
        style=doc.styles[name];style.font.name='Noto Sans KR';style.font.color.rgb=RGBColor(0,0,0);style.font.size=Pt(11 if name=='Normal' else 22 if name=='Title' else 14)
        style.paragraph_format.space_after=Pt(10)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.add_paragraph('환불 안내 구조 비교 실습','Title')
    doc.add_paragraph('교육용 합성 정책입니다. 실제 회사의 환불 기준이 아닙니다. 문단과 표의 조건을 원본과 대조하세요.')
    doc.add_paragraph('신청 기한은 7일입니다.')
    doc.add_paragraph('영수증을 함께 제출합니다.')
    doc.add_paragraph('처리 기준','Heading 1')
    table=doc.add_table(rows=3,cols=3);table.autofit=False
    for col,width in zip(table.columns,[2.4,2.5,1.0]):col.width=Inches(width)
    for rn,row in enumerate(table.rows):
        if rn==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cn,cell in enumerate(row.cells):
            cell.text=ROWS[rn][cn];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr();shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'DCE9E2' if rn==0 else 'FFFFFF');props.append(shading)
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:color'),'D9D9D9');border.set(qn('w:sz'),'6');borders.append(border)
            props.append(borders);margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                margin=OxmlElement('w:'+edge);margin.set(qn('w:w'),'120');margin.set(qn('w:type'),'dxa');margins.append(margin)
            props.append(margins)
            for p in cell.paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.CENTER if cn==2 else WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph('표의 금액만 보지 말고 해당 행과 열의 조건을 함께 확인하세요.')
    paragraphs=list(doc.paragraphs)+[p for row in table.rows for cell in row.cells for p in cell.paragraphs]
    for p in paragraphs:
        for run in p.runs:
            run.font.name='Noto Sans KR';run.font.color.rgb=RGBColor(0,0,0);fonts=run._element.get_or_add_rPr().rFonts
            for name in list(fonts.attrib):
                if name.endswith('Theme'):del fonts.attrib[name]
            for script in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+script),'Noto Sans KR')
    doc.core_properties.author='K-Edu synthetic teaching lab';doc.save(output/'policy.docx')
    (output/'policy.txt').write_text('교육용 합성 정책\n신청 기한은 7일입니다.\n영수증을 함께 제출합니다.\n미사용 | 기한 7일 | 부담금 0원\n사용 이력 | 기한 개별 검토 | 부담금 15000원\n',encoding='utf-8')
    (output/'gold.json').write_text(json.dumps(dict(synthetic=True,docx_table=ROWS,docx_header_rows=[1],docx_paragraph='신청 기한은 7일입니다.',txt_line2='신청 기한은 7일입니다.'),ensure_ascii=False,indent=2)+'\n')


if __name__=='__main__':
    p=argparse.ArgumentParser(__doc__);p.add_argument('--output',required=True);build(p.parse_args().output)
